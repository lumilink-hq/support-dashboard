# Builds docs/LumiLink-AI-Concierge-Scope.docx. Edit the content section, then run:
#   python scripts/build-ai-concierge-scope.py   (needs: pip install python-docx)
# Styling comes from the scope-brief skill; findings live in docs/ai-concierge-findings.md.
import os
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Twips
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH

INK = "111111"         # body text and rules
MUTED = "3A3A3A"       # subtitle
GREY = "5A5A5A"        # deck lines, meta
FAINT = "6A6A6A"       # sources
HDR_GREY = "7A7A7A"    # running header
FOOT_GREY = "9A9A9A"   # footer date
WHITE = "FFFFFF"
ZEBRA = "F2F2F2"       # alternating row fill
HAIRLINE = "BFBFBF"    # row separators
CONTENT_W = 10080      # twips, 7.0in inside 0.75in margins


def _rgb(h):
    return RGBColor.from_string(h)


def _el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn("w:" + k), str(v))
    return e


def _fmt(run, size, color=INK, bold=False, caps=False, track=None):
    """size in pt. track = letter spacing in twentieths of a point."""
    run.font.name = "Arial"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.all_caps = caps
    run.font.color.rgb = _rgb(color)
    if track:
        run._element.rPr.append(_el("w:spacing", val=track))
    return run


def _para(host, before=0, after=4, style=None, line=264, border=None):
    """host is the document or a table cell. border = (color, sz, space)."""
    p = host.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    if line:
        p._p.get_or_add_pPr().append(_el("w:spacing", line=line, lineRule="auto"))
    if border:
        pBdr = OxmlElement("w:pBdr")
        pBdr.append(_el("w:bottom", val="single", color=border[0],
                        sz=border[1], space=border[2]))
        p._p.get_or_add_pPr().append(pBdr)
    return p


def _field(paragraph, instr):
    r = paragraph.add_run()
    r._r.append(_el("w:fldChar", fldCharType="begin"))
    t = OxmlElement("w:instrText")
    t.set(qn("xml:space"), "preserve")
    t.text = instr
    r._r.append(t)
    r._r.append(_el("w:fldChar", fldCharType="end"))
    return r


def new_doc(running_head=None, footer_date=None):
    d = Document()
    for s in d.sections:
        s.page_width, s.page_height = Inches(8.5), Inches(11)
        s.left_margin = s.right_margin = Inches(0.75)
        s.top_margin = s.bottom_margin = Inches(0.75)
        s.header_distance = s.footer_distance = Twips(640)
    n = d.styles["Normal"]
    n.font.name = "Arial"
    n.font.size = Pt(9.5)
    n.font.color.rgb = _rgb(INK)
    if running_head:
        hp = d.sections[0].header.paragraphs[0]
        hp.paragraph_format.space_after = Pt(6)
        pBdr = OxmlElement("w:pBdr")
        pBdr.append(_el("w:bottom", val="single", color="D0D0D0", sz=4, space=4))
        hp._p.get_or_add_pPr().append(pBdr)
        _fmt(hp.add_run(running_head), 7.5, HDR_GREY, track=30)
    if footer_date:
        fp = d.sections[0].footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _fmt(fp.add_run(footer_date + "        "), 7.5, FOOT_GREY)
        _fmt(_field(fp, "PAGE"), 7.5, HDR_GREY)
    return d


def title(doc, eyebrow, headline, subtitle=None, meta=None):
    _fmt(_para(doc, 12, 2, line=None).add_run(eyebrow), 8.5, INK, True, True, 60)
    _fmt(_para(doc, 0, 2.5, line=None).add_run(headline), 22, INK, True)
    if subtitle:
        _fmt(_para(doc, 0, 8, line=None,
                   border=(INK, 12, 10)).add_run(subtitle), 10.5, MUTED)
    if meta:
        _fmt(_para(doc, 5, 10).add_run(meta), 8.5, GREY)


def h1(doc, text):
    p = _para(doc, 18, 5, line=None, border=(INK, 12, 6))
    p.paragraph_format.keep_with_next = True
    _fmt(p.add_run(text), 12, INK, True, True, 30)


def h2(doc, text):
    p = _para(doc, 10, 3, line=None)
    p.paragraph_format.keep_with_next = True
    _fmt(p.add_run(text), 10, INK, True, track=12)


def deck(doc, text):
    """One grey line under a heading, framing the table that follows."""
    p = _para(doc, 0, 5)
    p.paragraph_format.keep_with_next = True
    _fmt(p.add_run(text), 8.5, GREY)


def bullet(doc, segments):
    """segments: a string, or [(text, bold), ...] to bold the lead-in."""
    if isinstance(segments, str):
        segments = [(segments, False)]
    p = _para(doc, 2, 2, style="List Bullet")
    for text, bold in segments:
        _fmt(p.add_run(text), 9.5, INK, bold)


def lb(lead, rest):
    """Bold lead-in plus plain text, for bullet()."""
    return [(lead, True), (rest, False)]


def sources(doc, text):
    _fmt(_para(doc, 10, 6).add_run(text), 7.5, FAINT)


def _cell(c, text, size=8.5, bold=False, color=INK, caps=False, track=None,
          fill=None):
    tcPr = c._tc.get_or_add_tcPr()
    if fill:
        tcPr.append(_el("w:shd", val="clear", color="auto", fill=fill))
    mar = OxmlElement("w:tcMar")
    for side, w in (("top", 90), ("left", 110), ("bottom", 90), ("right", 110)):
        mar.append(_el("w:" + side, type="dxa", w=w))
    tcPr.append(mar)
    p = c.paragraphs[0]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    p._p.get_or_add_pPr().append(_el("w:spacing", line=264, lineRule="auto"))
    _fmt(p.add_run(str(text)), size, color, bold, caps, track)


def _grid(t, widths):
    """Fixed layout, hairline horizontals, no verticals, rows never split."""
    t.autofit = False
    tblPr = t._tbl.tblPr
    tblPr.append(_el("w:tblLayout", type="fixed"))
    tblPr.append(_el("w:tblW", type="dxa", w=CONTENT_W))
    b = OxmlElement("w:tblBorders")
    b.append(_el("w:top", val="single", color=HAIRLINE, sz=6))
    b.append(_el("w:bottom", val="single", color=HAIRLINE, sz=6))
    b.append(_el("w:insideH", val="single", color=HAIRLINE, sz=4))
    for side in ("left", "right", "insideV"):
        b.append(_el("w:" + side, val="none", color=WHITE, sz=0))
    tblPr.append(b)
    for gridCol, w in zip(t._tbl.find(qn("w:tblGrid")), widths):
        gridCol.set(qn("w:w"), str(int(CONTENT_W * w)))
    for row in t.rows:
        row._tr.get_or_add_trPr().append(_el("w:cantSplit", val="true"))
        for i, w in enumerate(widths):
            row.cells[i].width = Twips(int(CONTENT_W * w))


def _keep_together(t):
    """Hold a short table on one page: keepNext on every row but the last."""
    for row in t.rows[:-1]:
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.keep_with_next = True


def table(doc, headers, rows, widths, bold_last=False, keep_together=False):
    """widths: column fractions summing to 1.0. bold_last: bold a TOTAL row.
    keep_together: for a short table that must not break across pages."""
    assert abs(sum(widths) - 1.0) < 0.001, sum(widths)
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, h in enumerate(headers):
        _cell(t.rows[0].cells[i], h, 8.5, True, WHITE, True, 16, INK)
    t.rows[0]._tr.get_or_add_trPr().append(_el("w:tblHeader", val="true"))
    for r_i, row in enumerate(rows):
        last = bold_last and r_i == len(rows) - 1
        fill = ZEBRA if (last or r_i % 2 == 1) else None
        cells = t.add_row().cells
        for i, val in enumerate(row):
            _cell(cells[i], val, 8.5, last, INK, fill=fill)
    _grid(t, widths)
    if keep_together:
        _keep_together(t)
    _para(doc, 0, 4, line=None)
    return t


def kv_table(doc, rows, widths=(0.22, 0.78)):
    """Headerless summary block: bold label, zebra rows. rows = [(k, v), ...]"""
    assert abs(sum(widths) - 1.0) < 0.001, sum(widths)
    t = doc.add_table(rows=0, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for r_i, (k, v) in enumerate(rows):
        fill = ZEBRA if r_i % 2 == 1 else None
        cells = t.add_row().cells
        _cell(cells[0], k, 8.5, True, INK, fill=fill)
        _cell(cells[1], v, 8.5, False, INK, fill=fill)
    _grid(t, widths)
    _para(doc, 0, 4, line=None)
    return t


def callout(doc, heading, paragraphs):
    """One boxed aside with a heavy left rule. Use it once, for the one thing."""
    _para(doc, 0, 8, line=None)  # the box needs air above it
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False
    tblPr = t._tbl.tblPr
    tblPr.append(_el("w:tblLayout", type="fixed"))
    tblPr.append(_el("w:tblW", type="dxa", w=CONTENT_W))
    b = OxmlElement("w:tblBorders")
    b.append(_el("w:left", val="single", color=INK, sz=24))
    for side in ("top", "bottom", "right"):
        b.append(_el("w:" + side, val="single", color=INK, sz=6))
    for side in ("insideH", "insideV"):
        b.append(_el("w:" + side, val="none", color=WHITE, sz=0))
    tblPr.append(b)
    t._tbl.find(qn("w:tblGrid"))[0].set(qn("w:w"), str(CONTENT_W))
    c = t.rows[0].cells[0]
    c.width = Twips(CONTENT_W)
    _cell(c, heading, 8.5, True, INK, True, 20)
    for para in paragraphs:
        _fmt(_para(c, 1.5, 1.5).add_run(para), 9, INK)
    t.rows[0]._tr.get_or_add_trPr().append(_el("w:cantSplit", val="true"))
    _para(doc, 0, 4, line=None)
    return t



# ---------------------------------------------------------------- content

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
# Set OUT to build somewhere else (e.g. to diff against the hand-edited docx).
OUTPUT_PATH = os.environ.get("OUT") or os.path.join(ROOT, "docs", "LumiLink-AI-Concierge-Scope.docx")
DATE = "October 15, 2026"
d = new_doc(running_head="LUMILINK  ·  AI CONCIERGE SCOPE", footer_date=DATE)
title(d, "LumiLink", "AI Budtender")

h1(d, "Summary")
kv_table(d, [
    ("The ask", "A talking AI budtender on the website: trained on the store, guardrailed, "
                "a set personality and a realistic custom look."),
    ("What gets built", "A LumiLink add-on: our existing ElevenLabs agent (voice, 30+ languages, "
                        "knowledge base, product lookup) shown as a realistic AI human, streamed live by HeyGen LiveAvatar."),
    ("Effort", "35-51 hours for the LumiLink add-on, then 8-14 to launch it for PACKS; about 2 weeks."),
    ("Running cost", "About $0.18-0.22 per conversation minute, plus about $99/month for LiveAvatar."),
    ("Stack", "ElevenLabs Agents for brain and voice; HeyGen LiveAvatar (face only) for the avatar, launching on a stock AI human."),
])

h1(d, "Effort")
deck(d, "One developer, Claude-assisted. Hours include build and testing.")
rows = [
    ["LiveAvatar setup: stock AI human, latency and lip-sync tests", 3, 5],
    ["Website embed: secure session token, avatar player, mic, text fallback, mobile", 6, 9],
    ["Connect the ElevenLabs agent's audio to the avatar's lip sync", 3, 4],
    ["Persona and guardrail profile, plus a red-team test set", 7, 10],
    ["Store knowledge: reuse knowledge base and product lookup, add a recommendation tool", 2, 4],
    ["Language detection and testing in the top 5 customer languages", 3, 4],
    ["Dashboard: avatar and persona settings, kill switch, transcripts, minute caps", 4, 5],
    ["Add-on billing, entitlement and pricing page", 1, 2],
    ["Load, concurrency and launch QA", 6, 8],
]
lo = sum(r[1] for r in rows); hi = sum(r[2] for r in rows)
assert (lo, hi) == (35, 51), (lo, hi)
table(d, ["Module", "Hrs"],
      [[r[0], f"{r[1]}-{r[2]}"] for r in rows]
      + [["TOTAL, LumiLink add-on", f"{lo}-{hi}"]],
      [0.74, 0.26], bold_last=True, keep_together=True)
bullet(d, lb("Basis: ", "SEO was scoped at 173-210 hours and its code shipped in a week; Outbound phases 1, 2 and 4 took about 3 hours."))
bullet(d, lb("Ceiling: ", "the original estimate was 82-118 hours; budget against it if the mobile embed slips."))

h1(d, "Cost")
table(d, ["Line", "Type", "Amount"],
      [["ElevenLabs agent minutes plus LLM tokens", "Variable", "$0.085-0.125/min"],
       ["Avatar streaming (LiveAvatar, face only)", "Variable", "About $0.09/min"],
       ["LiveAvatar plan", "Recurring", "About $99/month"],
       ["Custom PACKS avatar (optional; launch on stock)", "One-time", "LiveAvatar quote, if wanted"],
       ["Total per conversation minute", "", "$0.18-0.22"],
       ["Example: 1,000 four-minute chats a month", "", "$700-860 plus $99"]],
      [0.45, 0.17, 0.38], bold_last=True, keep_together=True)

h1(d, "Rules and Risks")
bullet(d, lb("Cannabis: ", "no medical or effect claims, a 21+ gate before the avatar talks, counsel review before launch."))
bullet(d, lb("California: ", "disclose recording (two-party consent) and that the shopper is talking to an AI."))
bullet(d, lb("Health data: ", "shoppers will describe conditions; transcripts need retention limits and a privacy notice."))
bullet(d, lb("Cost runaway: ", "video costs double voice per minute; per-client minute caps ship day one."))

h1(d, "PACKS: On LumiLink or In-House")
deck(d, "The same budtender, built two ways, assuming one developer builds either. Detail: the PACKS In-House Build scope.")
table(d, ["", "On the LumiLink stack", "Built in-house at PACKS"],
      [["Total build hours", "43-65: add-on 35-51, PACKS launch 8-14", "45-66"],
       ["Monthly upkeep", "4-8 hrs; shared only once other clients join", "4-8 hrs"],
       ["Running cost", "Same usage plus a LumiLink add-on fee, not set", "Same usage plus about $120-205/month in plans"],
       ["Customer data", "PACKS transcripts sit in LumiLink's systems", "Stays with PACKS"],
       ["Owner", "LumiLink; PACKS sets persona, look and menu", "PACKS's one developer"],
       ["Comes out ahead when", "LumiLink sells it to other clients", "PACKS is the only user"]],
      [0.20, 0.38, 0.42], keep_together=True)

d.save(OUTPUT_PATH)

from docx import Document as D2
doc = D2(OUTPUT_PATH)
words = sum(len(p.text.split()) for p in doc.paragraphs)
for t in doc.tables:
    for r in t.rows:
        for c in r.cells:
            words += len(c.text.split())
print("words", words)
print("emdash", "\u2014" in "\n".join(p.text for p in doc.paragraphs))
